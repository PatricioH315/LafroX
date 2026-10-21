from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.shared import Inches
from docx.oxml.ns import qn

scripts={
34:[
'El plan de calidad convierte las normas en controles verificables. ISO 25010 organiza las características del producto, como desempeño, fiabilidad y seguridad; ISO 25012 revisa la calidad de los datos, incluyendo su completitud y consistencia. Para demostrarlo, cada característica tiene una métrica, un umbral y una prueba asociada. ISO/IEC/IEEE 29119 estructura cómo planificamos, ejecutamos y documentamos esas pruebas.',
'Complementamos ese marco con análisis de seguridad basados en OWASP y revisión de accesibilidad WCAG 2.2 AA, tanto automática como manual. Las exigencias sanitarias, tributarias y de operación también se convierten en escenarios de verificación. Para Puelche, la evidencia no será una lista de normas: serán informes, registros de ejecución y resultados que permitan comprobar el cumplimiento antes de aceptar cada entrega.'
],
35:[
'La calidad del producto se evalúa con cuatro objetivos y ocho características medibles. El 100 % de los requisitos ofertados debe tener una prueba aprobada antes de su hito, y la marcha blanca no puede cerrar con defectos críticos o altos abiertos. En desempeño usamos el percentil 95, no el promedio: al menos el 95 % de las operaciones debe cumplir el umbral de preparación de un segundo o preventa de un segundo y medio, bajo la carga definida.',
'En fiabilidad distinguimos disponibilidad mensual y continuidad del despacho: la transacción crítica debe alcanzar al menos 99,9 %, pero entre 05:30 y 07:00 exigimos cero minutos de caída. Además, ensayamos RTO de hasta cuatro horas para recuperar servicios críticos y RPO de hasta quince minutos para los datos. La mantenibilidad se controla con cobertura unitaria de al menos 70 % en negocio y 80 % global; a Producción se promueve el mismo artefacto probado.'
],
36:[
'Las siete puertas de calidad aplican controles progresivos al cambio. G0 y G1 revisan pruebas básicas y aprobación por pares. G2 ejecuta automáticamente las pruebas y métricas de código; G3 bloquea la promoción a QA si hay vulnerabilidades críticas o altas o regresión de desempeño. G4 verifica regresión, análisis dinámico de seguridad y accesibilidad antes de promover a preproducción.',
'G5 habilita cada etapa en Producción al iniciar su marcha blanca: exige certificación aceptada, pruebas aplicables, conciliación, capacitación y reversión ensayada. Después, G6 observa el cambio mediante despliegue gradual o canario y revierte si degrada el servicio. Estas puertas no sustituyen la aceptación de Puelche. Para cerrar la marcha blanca deben cumplirse las seis condiciones: sin incidentes graves abiertos atribuibles a la solución, volumen real completo durante las últimas cuatro semanas, disponibilidad y respuesta exigidas, conciliación, usuarios de la ola certificados y acta firmada.'
],
37:[
'El plan distribuye las pruebas entre desarrollo, QA, preproducción, terreno y operación. Las unitarias revisan lógica aislada; las de integración comprueban las quince conexiones; la regresión verifica que un cambio no rompa funciones existentes. La aceptación incorpora a los usuarios que ejecutarán las tareas reales. Para carga utilizamos 21,99 transacciones por segundo, equivalentes a 1,5 veces el máximo de referencia; con estrés aumentamos la demanda para comprobar degradación controlada sin pérdida de transacciones.',
'La continuidad se prueba con sitios veinticuatro horas sin enlace y terminales catorce horas sin señal, además de condiciones físicas como frío con guantes. El catálogo inicial tiene 1.316 casos, con 80 % de automatización prevista. Las 4,4 horas corresponden únicamente a esa parte automática en dos ejecutores paralelos; las pruebas manuales se programan aparte. Durante la marcha blanca medimos diariamente los indicadores con operación real.'
],
38:[
'La trazabilidad conecta el resultado esperado con el requisito, el módulo, el trabajo programado y la prueba que lo demuestra. En el ejemplo, identificar afectados por un lote en menos de dos horas se vincula con RF-09.01, el módulo M9 y un simulacro de retiro. La solicitud de fusión registra el requisito y el digest de la imagen identifica exactamente la versión desplegada; así podemos relacionar el resultado de la prueba con lo que recibió Puelche.',
'Si aparece una falla, aplicamos PDCA: identificamos causa y responsable, corregimos, verificamos con el mismo criterio y cerramos o ajustamos. Pareto ayuda a priorizar causas y los diagramas de control a detectar variaciones anormales. El plan asigna 11.136 horas-hombre a prevención y evaluación, aproximadamente 5,1 % de la línea base. Es conformidad programada, no el costo integral de todas las fallas posibles.'
],
50:[
'Las cinco innovaciones agregan capacidades al alcance obligatorio y tienen indicadores de validación. La primera estima el riesgo de vencimiento en el local usando entregas, devoluciones y ritmo de reposición; el preventista confirma y actúa. La meta es reducir 25 % el vencido local frente a la referencia. La segunda reconstruye incidentes con evidencia sanitizada, reproduce su causa en pruebas y conserva el escenario como regresión; busca reducir 30 % la mediana de horas hasta reproducir la causa.',
'La tercera estima vida útil con temperatura y tiempo mediante un modelo determinista por familia. Se valida con al menos sesenta lotes por familia: al menos 80 % debe quedar dentro del margen de ±15 % de la vida nominal respecto de la evaluación de Calidad. Nunca amplía la fecha impresa ni elimina bloqueos. La cuarta vincula un tramo variable del pago al costo comparable por entrega: meta de al menos 5 % de reducción con OTIF de al menos 95 %, usando una base y ajustes definidos. La quinta entrega información al almacenero en el terminal sin señal o mediante copia solicitada; la comparación con otros locales requiere aprobación legal. Los beneficios se verifican: no se dan por demostrados.'
]
}
root=Path.cwd()
src=root/'Presentaciones'/'Guia breve para el cliente SD9 y SD13.docx'
out=root/'Presentaciones'/'Guia para el cliente SD9 y SD13 con guion tecnico.docx'
d=Document(src)
table_before=[[c.text for row in t.rows for c in row.cells] for t in d.tables]
heads=[p for p in d.paragraphs if p.text=='Guion breve']
assert len(heads)==6
for n,h in zip([34,35,36,37,38,50],heads):
 h.text='Guion para exponer'
 para=Paragraph(h._p.getnext(),h._parent)
 para.text=scripts[n][0]
 para.paragraph_format.keep_with_next=False
 new=OxmlElement('w:p');para._p.addnext(new)
 p=Paragraph(new,para._parent);p.style=para.style;p.text=scripts[n][1]
for section in d.sections:
 section.top_margin=Inches(.5);section.bottom_margin=Inches(.5)
for table in d.tables:
 for row in table.rows:
  for cell in row.cells:
   mar=cell._tc.get_or_add_tcPr().find(qn('w:tcMar'))
   if mar is not None:
    for side in ['top','bottom']:
     el=mar.find(qn('w:'+side))
     if el is not None:el.set(qn('w:w'),'45')
for p in list(d.paragraphs):
 if not p.text and p._p.findall('.//'+qn('w:br')):
  p._p.getparent().remove(p._p)
titles=[p for p in d.paragraphs if p.style.name=='Title']
for i,p in enumerate(titles):p.paragraph_format.page_break_before=(i>0)
assert table_before==[[c.text for row in t.rows for c in row.cells] for t in d.tables]
d.core_properties.title='Guía para el cliente con guion técnico de calidad e innovaciones'
d.save(out)
print(out)
print('Tablas conservadas',len(d.tables))
for n,parts in scripts.items():print('Diapositiva',n,'palabras de guion',sum(len(x.split()) for x in parts))
