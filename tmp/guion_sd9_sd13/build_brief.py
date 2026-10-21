from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH

slides=[
(34,'Normas y leyes',[
('ISO 25010','Es una pauta para revisar si el sistema funciona bien, es rápido, seguro y fácil de mantener. Cada aspecto tiene una medida.'),
('ISO 25012','Es una pauta para comprobar que los datos estén completos, sean correctos y no se contradigan entre sistemas.'),
('ISO/IEC/IEEE 29119','Ordena las pruebas: qué se revisa, cómo se ejecuta y qué evidencia queda. Sus tres niveles organizan el proceso de pruebas.'),
('PMBOK capítulo 8','Es una guía para planificar la calidad, revisar cómo trabajamos y comprobar que las entregas cumplen lo acordado.'),
('WCAG 2.2 AA','Son criterios para que personas con distintas capacidades puedan usar el sistema. Se revisan con herramientas y personas, no solo automáticamente.'),
('OWASP','Son referencias para detectar fallas de seguridad: revisamos tanto el código como la aplicación funcionando.'),
('ISO 9001','Es un marco para trabajar con procesos de calidad documentados y auditables. No significa que el programa nunca pueda fallar.'),
('CMMI nivel 3','Es una evaluación de cómo organiza la empresa su desarrollo: utiliza procesos definidos y repetibles, no depende de improvisar cada proyecto.'),
('Decreto 977 y frío','Es la referencia sanitaria de la lámina. El sistema registra temperaturas y ayuda a controlar el frío; −18 °C se contextualiza en congelados.'),
('SII y DTE','DTE es un documento tributario electrónico. Verificamos el intercambio de la guía electrónica y la constancia de recepción, sin reemplazar al emisor autorizado.'),
('Artículo 25 bis','En esta propuesta significa que las rutas deben considerar restricciones de jornada y descansos del conductor, no solo el camino más corto.'),
('Ley 21.719','Es la referencia de protección de datos personales. Se traduce en limitar accesos y evitar exponer información de personas.'),
('GS1 y EDI','GS1 permite identificar productos, lugares y unidades logísticas con códigos comunes. EDI permite intercambiar documentos entre empresas sin digitarlos nuevamente.'),
('GTIN, GLN y SSCC','GTIN identifica el producto; GLN, una entidad o ubicación; SSCC, una unidad logística, como un pallet identificado individualmente.')
],
'No basta con mencionar normas. Cada una nos ayuda a revisar algo concreto: el funcionamiento del sistema, la calidad de los datos, la seguridad o la forma de probar. Lo importante para Puelche es que ese cumplimiento se demuestre con resultados y registros que pueda revisar.',
'SD9 §§9.1–9.2; T13; diapositiva 34.'),
(35,'Calidad del producto',[
('4 objetivos','Probar lo ofertado, cerrar sin defectos graves pendientes, proteger la operación y demostrar los resultados de negocio en el momento previsto.'),
('100 % de requisitos probados','Cada requisito ofertado debe tener una prueba aprobada antes del hito que le corresponde. No basta con decir que está desarrollado.'),
('0 defectos críticos o altos','No se cierra la marcha blanca con problemas graves pendientes. No significa que jamás pueda aparecer una falla.'),
('0 minutos entre 05:30 y 07:00','Es la ventana de despacho que no puede interrumpirse. Un buen porcentaje mensual no compensa una caída en ese horario.'),
('16 de 16 resultados','El éxito incluye los dieciséis resultados del caso, no solo entregar software. Cada resultado se verifica cuando corresponde.'),
('p95 y tiempos de respuesta','Al menos el 95 % de las operaciones medidas debe cumplir el tiempo: preparación ≤ 1 segundo y preventa ≤ 1,5 segundos.'),
('15 integraciones','Son conexiones con otros componentes o sistemas. Se prueba que las quince intercambien información correctamente.'),
('Usabilidad y guantes a −22 °C','El sistema debe poder usarse en el trabajo real, incluso en frío con guantes; no solo verse bien en una oficina.'),
('Disponibilidad ≥ 99,9 %','Es la proporción del mes en que la transacción crítica debe estar disponible. Se controla además la ventana protegida del despacho.'),
('RTO ≤ 4 horas','Es el tiempo objetivo máximo para recuperar los servicios críticos ante el desastre previsto. No autoriza cuatro horas de caída en despacho.'),
('RPO ≤ 15 minutos','Es el máximo de información reciente que podría perderse al recuperar el sistema, expresado como tiempo. No es cuánto demora restaurarlo.'),
('0 hallazgos de seguridad graves','No se permite avanzar con vulnerabilidades críticas o altas pendientes. No equivale a afirmar que el riesgo es cero.'),
('Cobertura 70 % y 80 %','Mide cuánto código ejercitan las pruebas: al menos 70 % de la lógica de negocio y 80 % global. No garantiza ausencia de errores.'),
('Mismo artefacto','Se lleva a Producción la misma versión de software que se probó. Así evitamos aprobar una versión y entregar otra.')
],
'Aquí calidad significa cumplir medidas concretas. Probamos cada requisito, verificamos tiempos de respuesta y protegemos el despacho. También comprobamos conexiones, seguridad y uso en terreno. La recuperación se ensaya: RTO indica cuánto tardamos en restablecer el servicio y RPO cuánto dato reciente podría perderse.',
'SD9 Tablas 9.1, 9.3 y 9.4; Anexo 9.A; T17.'),
(36,'Calidad que bloquea',[
('7 puertas de calidad','Son siete puntos de control. Si un cambio no cumple una condición obligatoria, se detiene hasta corregirlo.'),
('G0 cambio local','El desarrollador comprueba formato y pruebas básicas antes de compartir su cambio.'),
('G1 revisión por pares','Otra persona revisa el código antes de incorporarlo. No depende únicamente de quien lo escribió.'),
('G2 integración continua','Un proceso automático ejecuta pruebas y revisa cobertura, complejidad y acuerdos de intercambio entre componentes.'),
('G3 promoción a QA','Antes de pasar al ambiente de pruebas, se bloquean vulnerabilidades críticas o altas y regresiones de desempeño.'),
('G4 preproducción','Es el ambiente previo al uso real. Se comprueba que el cambio no rompa funciones anteriores y se revisan seguridad y accesibilidad.'),
('DAST y regresión','DAST busca fallas de seguridad con la aplicación funcionando. Regresión vuelve a probar lo existente después de un cambio.'),
('G5 entrada a Producción','Permite iniciar la marcha blanca solo con pruebas aprobadas, datos conciliados, usuarios capacitados y regreso a la versión anterior ensayado.'),
('G6 producción supervisada','Se observa el cambio mientras funciona con usuarios reales y se revierte si empeora el servicio.'),
('Canario','Es probar una versión con una parte limitada de la operación antes de extenderla, para reducir el impacto si falla.'),
('Presupuesto de error','Es el margen de fallas compatible con la meta del servicio. Al agotarse, se restringen nuevos cambios no urgentes.'),
('6 condiciones para aceptar','Deben cumplirse todas: sin incidentes graves abiertos; 100 % del volumen real de la ola durante las últimas 4 semanas; disponibilidad y respuesta exigidas; sin diferencias no explicadas; usuarios de la ola certificados; acta firmada.'),
('Acta de la Contraparte','Es la aceptación formal de Puelche. Pasar nuestros controles internos no nos permite declarar aceptado el hito por nuestra cuenta.')
],
'Las puertas impiden que un cambio avance solo porque tenemos prisa. Primero revisamos el código y las pruebas; luego verificamos su comportamiento antes y durante el uso real. Pero la aceptación final corresponde a Puelche: requiere cumplir las seis condiciones y firmar el acta. Ninguna condición compensa a otra.',
'SD9 §9.2.1; Anexo 9.B; T17 §4.3.'),
(37,'Pruebas del código al terreno',[
('12 tipos de prueba','Son distintas formas de revisar el sistema. Cada prueba debe dejar un resultado medible o una validación documentada.'),
('Unitarias y de componente','Comprueban por separado partes del programa para encontrar errores antes de conectarlas con el resto.'),
('Contratos de API','Verifican que dos componentes sigan entendiéndose: qué información se envía, en qué formato y qué respuesta se espera.'),
('Integración y sistema','Integración comprueba las conexiones; sistema verifica que el conjunto permita completar las tareas del negocio.'),
('Regresión, DAST y accesibilidad','Revisamos que no se rompa lo anterior, que no haya fallas graves de seguridad y que las personas puedan usar la aplicación.'),
('Aceptación de usuario','Personas que realizarán el trabajo prueban sus tareas para comprobar que la solución responde a lo que necesitan.'),
('Carga 1,5× y estrés','Se prueba con 50 % más demanda que el máximo previsto; después se aumenta más para observar cómo responde al saturarse.'),
('21,99 TPS','Son 21,99 transacciones por segundo de carga de prueba. Resulta de multiplicar el máximo de 14,66 por 1,5.'),
('Resiliencia y recuperación','Resiliencia es resistir fallas y seguir operando; recuperación es restablecer el servicio y sus datos después del desastre previsto.'),
('24 horas y 14 horas','Se prueban sitios durante 24 horas sin enlace y terminales durante 14 horas sin señal. Son escenarios distintos.'),
('Pruebas de terreno','Se comprueba uso con guantes a −22 °C, lluvia, calor superior a 34 °C y una mano: condiciones reales, no simuladas solo en pantalla.'),
('Migración, despliegue e intrusión','Revisamos el traslado correcto de datos, la instalación y reversión del cambio y la seguridad mediante un especialista externo autorizado.'),
('1.316 casos y 80 % automático','Es el catálogo inicial previsto; unos 1.053 casos se automatizan. Los demás necesitan intervención humana.'),
('4,4 horas','Es la duración estimada de la parte automatizada con dos ejecutores paralelos. No incluye todo el trabajo manual o de terreno.'),
('Datos y marcha blanca','En pruebas usamos datos sintéticos o anonimizados. En marcha blanca medimos diariamente el funcionamiento con operación real.')
],
'No probamos únicamente que una pantalla abra. Revisamos partes del programa, conexiones, seguridad y tareas reales. Lo sometemos a más carga y a condiciones de terreno. La parte automática puede ejecutarse de noche; las pruebas humanas se programan aparte. Durante la marcha blanca comprobamos diariamente cómo funciona con el volumen real.',
'SD9 §9.2.3; T13, tipos de prueba y §8.'),
(38,'Trazabilidad y mejora',[
('Trazabilidad de extremo a extremo','Es poder seguir una necesidad desde lo que pidió Puelche hasta el cambio construido, la prueba realizada y la aceptación.'),
('Retiro en menos de 2 horas','El compromiso es identificar a los clientes afectados por un lote en menos de dos horas, no retirar físicamente toda la mercadería.'),
('RF-09.01 y M9','RF-09.01 es el código del requisito de trazabilidad. M9 es el módulo de calidad y trazabilidad que participa en su cumplimiento.'),
('EDT y paquete 3.4.4','La EDT divide el proyecto en trabajos controlables. El código 3.4.4 ubica el trabajo de cadena de frío y retiro.'),
('Simulacro de retiro','Es un ensayo con un lote para comprobar que podemos identificar afectados y demostrar el tiempo comprometido.'),
('Fusión e ID del requisito','Al incorporar un cambio de código registramos qué requisito atiende. Así podemos encontrar quién cambió qué y para qué.'),
('Digest de imagen firmada','Es una huella que identifica la versión exacta del software; la firma ayuda a comprobar su origen e integridad.'),
('Acta H7','Es la aceptación del cierre de marcha blanca de la Etapa 1. No equivale a aceptar anticipadamente resultados que se medirán después.'),
('PDCA','Es un ciclo de mejora: planificar la corrección, hacerla, verificar si funcionó y actuar para cerrar o ajustar.'),
('Diagramas de control','Muestran cómo cambia una medida en el tiempo y ayudan a detectar comportamientos anormales del proceso.'),
('Pareto y causa raíz','Pareto ayuda a priorizar los problemas que más aportan al resultado. Causa raíz busca por qué ocurren, no solo su síntoma.'),
('11.136 HH y 5,1 %','Es el esfuerzo programado de prevención y evaluación: 5,1 % de las horas del proyecto. No incluye por sí solo todas las fallas futuras.'),
('37,9 % y 62,1 %','De esas horas, 37,9 % se dedica a prevenir problemas y 62,1 % a comprobar el cumplimiento mediante evaluaciones y pruebas.')
],
'La trazabilidad nos permite demostrar qué hicimos para cumplir cada compromiso. Podemos seguir el requisito, el cambio, la versión probada y su aceptación. Si falla, buscamos la causa, corregimos y volvemos a medir. El esfuerzo de prevención y evaluación está programado; no esperamos al final para preocuparnos por la calidad.',
'SD9 §§9.1.6, 9.2.5 y 9.2.6; Anexos 9.C y 9.D; T17.'),
(50,'Cinco innovaciones',[
('1 Vencimiento posentrega','Avisamos al preventista qué productos podrían vencer en el local para que confirme y actúe. Meta: reducir 25 % el vencido local frente a la referencia, no toda la merma.'),
('2 Reproducción de incidentes','Reconstruimos una falla de terreno en pruebas para entender su causa y verificar la corrección. Meta: reducir 30 % la mediana de horas hasta reproducir la causa.'),
('3 Vida útil por historia térmica','Usamos temperatura y tiempo para estimar vida útil restante. No amplía la fecha impresa ni elimina bloqueos. Se valida por familia: ≥80 % de lotes dentro del margen de ±15 % de vida nominal.'),
('4 Pago por costo de servir','Parte del pago depende de reducir el costo comparable por entrega al menos 5 %, manteniendo entregas a tiempo y completas —OTIF— de al menos 95 %. Ahorrar empeorando el servicio no cuenta.'),
('5 Hoja del almacenero','Devolvemos información útil de sus compras y productos por vencer, en el terminal del preventista sin internet. Puede solicitar copia para entrega posterior; la comparación con otros locales necesita aprobación legal.'),
('Qué tienen de innovación','Agregan capacidades al mínimo exigido. Sus beneficios son metas que debemos validar, no resultados que demos por demostrados.')
],
'Proponemos cinco mejoras: anticipar vencimientos en el local, reproducir fallas de terreno, estimar vida útil por temperatura, vincular parte del pago al costo de servir y devolver información al almacenero. Cada una tiene una forma de comprobar su aporte. No prometemos beneficios sin medirlos ni permitimos que una innovación debilite controles sanitarios o el servicio.',
'SD13 §§13.1–13.5; Anexos 13.A–13.C; T19.')
]

d=Document();sec=d.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.6);sec.bottom_margin=Inches(.6)
sec.left_margin=Inches(.65);sec.right_margin=Inches(.65);sec.footer_distance=Inches(.25)
for style in d.styles:
 if style.type==1:
  style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
for b in list(d.styles.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
d.styles['Normal'].font.size=Pt(10.5)
d.styles['Normal'].paragraph_format.space_after=Pt(4)
d.styles['Normal'].paragraph_format.line_spacing=1.02
d.styles['Title'].font.size=Pt(20)
d.styles['Title'].paragraph_format.space_after=Pt(3)
d.styles['Heading 2'].font.size=Pt(12)
d.styles['Heading 2'].paragraph_format.space_before=Pt(8)
d.styles['Heading 2'].paragraph_format.space_after=Pt(4)
d.core_properties.title='Guía breve para responder al cliente sobre calidad e innovaciones'
d.core_properties.author='LafroX'
for idx,(n,title,rows,script,source) in enumerate(slides):
 if idx:d.add_page_break()
 d.add_paragraph(f'Diapositiva {n}',style='Title')
 h=d.add_paragraph(title);h.paragraph_format.space_after=Pt(7)
 for r in h.runs:r.bold=True;r.font.size=Pt(13)
 if idx==0:
  p=d.add_paragraph('Numeración por posición real en PowerPoint. Respuestas breves para explicarlas al cliente.');p.paragraph_format.space_after=Pt(7)
 t=d.add_table(rows=1,cols=2);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 t.columns[0].width=Inches(1.78);t.columns[1].width=Inches(5.42)
 for cell,width in zip(t.rows[0].cells,[1.78,5.42]):cell.width=Inches(width)
 t.rows[0].cells[0].text='Concepto de la diapositiva';t.rows[0].cells[1].text='Qué es y cómo explicarlo al cliente'
 trpr=t.rows[0]._tr.get_or_add_trPr();rep=OxmlElement('w:tblHeader');trpr.append(rep)
 for label,answer in rows:
  cells=t.add_row().cells
  cells[0].width=Inches(1.78);cells[1].width=Inches(5.42)
  cells[0].text=label;cells[1].text=answer
 for ri,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
  for ci,c in enumerate(row.cells):
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=c._tc.get_or_add_tcPr()
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   tcpr.append(borders)
   margin=OxmlElement('w:tcMar')
   for side,value in [('top','65'),('bottom','65'),('left','85'),('right','85')]:
    el=OxmlElement('w:'+side);el.set(qn('w:w'),value);el.set(qn('w:type'),'dxa');margin.append(el)
   tcpr.append(margin)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7E6E6');tcpr.append(sh)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
    for r in p.runs:
     r.font.size=Pt(10.5);r.font.bold=(ri==0 or ci==0)
 d.add_heading('Guion breve',level=2)
 d.add_paragraph(script)
 p=d.add_paragraph('Base documental: '+source);p.paragraph_format.space_before=Pt(4)
 for r in p.runs:r.font.size=Pt(8)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
f.add_run('LafroX   |   ').font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
out=Path('Presentaciones/Guia breve para el cliente SD9 y SD13.docx').resolve()
d.save(out)
print(out)
print('Tablas:',len(d.tables),'Filas de contenido:',sum(len(t.rows)-1 for t in d.tables))
