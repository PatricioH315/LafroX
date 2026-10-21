from pathlib import Path
import json, zipfile, xml.etree.ElementTree as ET
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path.cwd()
OUT=ROOT/'Presentaciones'/'Diccionario y guion SD9 y SD13 v10 alterna.docx'
PPT=ROOT/'Presentaciones'/'LafroX - Presentación Informe 2 v10 alterna.pptx'
BASE=Path(r'C:\Users\henri\Downloads\Informe_2_LafroX\Informe_2_LafroX')
sections=[]
def slide(n,title,intro):
    s={'n':n,'title':title,'intro':intro,'entries':[],'script':[],'notes':[]}
    sections.append(s); return s
def entry(s,title,meaning,limit,source):
    s['entries'].append((title,meaning,limit,source))

s=slide(34,'Normas y leyes','Idea central: cada estándar citado debe traducirse en un control y una evidencia; nombrarlo no acredita su cumplimiento.')
entry(s,'Norma y forma de acreditarla','Acreditar significa mostrar el mecanismo aplicado y el registro que demuestra su resultado: métricas, informes, auditorías o actas. El SD9 conecta estándares, controles, pruebas y aceptación.','Es una forma de trabajo propuesta, no prueba de que el proyecto ya esté certificado o ejecutado.','SD9 §9.1.2, Tabla 9.2 y Figura 9.1.')
entry(s,'ISO/IEC 25010 y ocho características con umbral','Modelo de calidad del producto: funcionalidad, desempeño, compatibilidad, usabilidad, fiabilidad, seguridad, mantenibilidad y portabilidad. Cada característica se mide con criterios del proyecto.','Los umbrales numéricos no los fija por sí sola la norma: provienen de las Bases, el Caso y compromisos de LafroX. El plan usa las ocho características y nombres exigidos en las Bases.','SD9 §9.1.3 y Tabla 9.3; A9, Anexo 9.A.')
entry(s,'ISO/IEC 25012 y datos medidos y conciliados','Ordena la calidad de los datos: completitud, exactitud, consistencia y actualidad. Conciliar consiste en comparar registros o totales entre origen y destino y explicar sus diferencias.','No basta con que la aplicación responda: datos incompletos o incoherentes pueden invalidar sus resultados. Debe conservarse evidencia de la medición.','SD9 Tabla 9.2; A9, Anexo 9.A.9.')
entry(s,'ISO/IEC 29119 y proceso de pruebas en tres niveles','El nombre completo usado en el informe es ISO/IEC/IEEE 29119. Los tres niveles de proceso son organizacional, gestión de pruebas y proceso dinámico de pruebas.','No son tres ambientes ni tres tipos de prueba. En el guion se utiliza el nombre completo de la norma aunque la diapositiva lo abrevia.','SD9 §9.2.3; T13 §2.')
entry(s,'PMBOK capítulo 8 y planificar gestionar y controlar','Planificar establece objetivos y criterios; gestionar comprueba que el proceso pueda producir calidad; controlar verifica el resultado de los entregables. Incluye seguimiento del costo de la calidad.','Aseguramiento del proceso, control del producto y aceptación del cliente son funciones diferentes.','SD9 §§9.1, 9.1.1, 9.1.6 y 9.2.6.')
entry(s,'WCAG 2.2 AA y herramienta más revisión manual','La accesibilidad se verifica con herramientas y con personas: navegación por teclado, lector de pantalla y otras comprobaciones que la automatización no resuelve por completo. AA es el nivel de conformidad declarado.','Pasar una herramienta automática no equivale a demostrar toda la accesibilidad. La facilidad de uso con guantes también requiere pruebas operacionales específicas.','SD9 Tabla 9.2 y §9.2.2; T13, pruebas de accesibilidad y perfil operacional.')
entry(s,'OWASP y análisis estático y dinámico','El análisis estático revisa código y dependencias sin ejecutar la aplicación; el dinámico observa el software en ejecución. OWASP aporta referentes para controles de seguridad y madurez.','No equivale a una certificación universal de ausencia de vulnerabilidades. Los hallazgos críticos o altos bloquean las promociones previstas en el plan.','SD9 §§9.1.4 y 9.2.2; A9, Anexo 9.B.')
entry(s,'ISO 9001 y marco de auditoría','Se utiliza el sistema de gestión de calidad de LafroX para auditar procesos y registros del proyecto.','La certificación del sistema de gestión de una empresa no certifica automáticamente cada versión del producto que desarrolla.','SD9 §9.1.4.')
entry(s,'CMMI DEV N3 y madurez del proceso','El nivel 3 corresponde a la madurez de procesos definidos que el informe declara como referencia organizacional para desarrollo y pruebas.','CMMI se acredita mediante evaluación; no conviene llamarlo certificación del producto ni confundir madurez del proceso con ausencia de defectos.','SD9 §9.1.4; diapositiva 34.')
entry(s,'Decreto 977 y temperatura de menos dieciocho grados con retiro expedito','La lámina vincula el marco sanitario con el control de frío, el registro de temperatura y la capacidad de identificar productos afectados. El límite mostrado se interpreta en el contexto de congelados de la propuesta.','No generalizar ≤ −18 °C a todos los alimentos ni reemplazar las reglas por familia, excursión y decisión de Calidad por una sola cifra. Es una síntesis de la diapositiva, no una revisión jurídica del reglamento.','Diapositiva 34; SD9, introducción y §9.1.3; S13 §13.3, límites de la estimación térmica.')
entry(s,'SII DTE y guía electrónica con acuse de recibo','DTE significa documento tributario electrónico. El punto de la lámina es verificar el intercambio y la evidencia documental asociados al despacho y su recepción.','La prueba informática no sustituye la emisión válida en el sistema autorizado ni implica que LafroX reemplace al ERP como emisor tributario.','Diapositiva 34; SD9 §9.2.3 y T13, pruebas de integración; T17 §5, resultados del caso.')
entry(s,'Artículo 25 bis y jornada de conductores en el ruteo','El ruteo debe incorporar las restricciones de jornada declaradas en la propuesta, en lugar de optimizar rutas suponiendo tiempo ilimitado de conducción.','La afirmación describe un control funcional a probar; esta guía no desarrolla una interpretación legal de turnos, descansos o excepciones.','Diapositiva 34; SD9 §9.2.5 y A9, Anexo 9.D, trazabilidad de resultados.')
entry(s,'Ley 21 719 y datos personales desde diciembre de 2026','La propuesta incorpora protección de datos en accesos, pruebas, paquetes de incidentes y entrega de información al almacenero. La diapositiva sitúa su referencia temporal en diciembre de 2026.','No presentar la mera mención de la ley como prueba de conformidad. Son necesarios controles verificables; el bloque comparativo de INN-05 tiene una aprobación legal específica.','Diapositiva 34; SD9 §9.2.4; S13 §§13.2.4 y 13.5.4; A13, Anexo 13.A.')
entry(s,'GS1 EDI y GTIN GLN SSCC','GS1 aporta identificadores: GTIN para artículos comerciales, GLN para entidades o ubicaciones y SSCC para unidades logísticas. EDI es el intercambio electrónico estructurado de documentos.','La columna de la lámina agrupa leyes, reglas y estándares; GS1 y EDI no son leyes. Identificar una unidad no sustituye registrar y relacionar sus movimientos.','Diapositiva 34; SD9, integración y trazabilidad; T13, contratos e integraciones; A13, contratos de lote y GTIN.')
s['script']=[
'Una vez definido cómo ejecutaremos el proyecto, falta responder cómo demostraremos que lo entregado cumple. Nuestro plan no consiste solamente en mencionar normas: cada una se traduce en un control y una evidencia que Puelche podrá revisar.',
'ISO 25010 organiza la calidad del producto e ISO 25012 la de los datos. Para las pruebas aplicamos ISO/IEC/IEEE 29119, con procesos de planificación, ejecución y documentación. Complementamos esos controles con accesibilidad, seguridad y auditoría.',
'Las exigencias operacionales también deben probarse: cadena de frío, documentación del despacho, restricciones de conducción y protección de datos. Los identificadores GS1 y el intercambio EDI permiten reconocer unidades y documentos de forma consistente.',
'La idea es que cada exigencia termine en evidencia, no en una declaración. En la siguiente diapositiva veremos los resultados concretos con los que juzgaremos la calidad.'
]
s['notes']=['No leer la columna de leyes como una lista de obligaciones jurídicas exhaustiva. Explicar su aplicación en la solución.','No llamar certificación de producto a ISO 9001 ni a la evaluación de madurez CMMI.']

s=slide(35,'Calidad del producto','Idea central: la calidad se evalúa con cuatro objetivos y con métricas específicas para ocho características del producto.')
entry(s,'Cuatro objetivos medibles','Son: entregar lo ofertado con pruebas aprobadas, cerrar sin defectos graves, sostener la operación y demostrar los dieciséis resultados del caso.','Los resultados del negocio se aceptan en su momento de medición; no se dan por obtenidos al aprobar una prueba de laboratorio.','SD9 §9.1.1 y Tabla 9.1; T17 §5.')
entry(s,'Cien por ciento de requisitos con prueba aprobada y funcionalidad','Cada requisito ofertado debe relacionarse con una prueba que lo verifique y aprobarse antes del hito al que corresponde. La lámina lo repite como objetivo y como característica de funcionalidad.','No significa que el 100 % de todas las exigencias de las Bases esté declarado Cumple. El T13 conserva visibles las brechas no ofertadas como cumplimiento.','SD9 Tablas 9.1 y 9.3; T13 §§1, 6 y 8.2.')
entry(s,'Cero defectos críticos o altos al cierre','La condición se refiere a defectos abiertos de esas severidades al cerrar cada marcha blanca.','No promete que nunca aparezca un defecto ni que no existan observaciones menores. Lo bloqueante es cerrar con defectos graves pendientes.','SD9 Tabla 9.1; A9, métrica 9A-02; T17 §§4.3 y 5.')
entry(s,'Cero minutos de caída entre 05 30 y 07 00','Es la ventana protegida del despacho. Se mide la indisponibilidad durante esos noventa minutos.','Es una condición adicional al porcentaje mensual de disponibilidad; no puede compensarse una caída en esa ventana con buen desempeño el resto del mes.','SD9 Tablas 9.1 y 9.3; T17, tabla de condiciones del art. 17.3, p. 12.')
entry(s,'Dieciséis de dieciséis resultados del Caso','El éxito incluye resultados verificables del negocio: trazabilidad, operación, integración y otros compromisos del catálogo del caso.','Tienen distintas ventanas y métodos de aceptación; no afirmar que los dieciséis se certifican todos en H5 o H10.','SD9 §9.1.1 y A9, Anexo 9.D; T17 §5.')
entry(s,'Desempeño p95 preparación de un segundo y preventa de un segundo y medio','p95 es el tiempo por debajo del cual queda al menos el 95 % de las observaciones. Los límites se verifican bajo la carga y concurrencia del plan.','No es tiempo promedio ni un segundo para cualquier operación. El informe agrega entrega y consulta de stock y crédito con umbrales propios.','SD9 §9.1.3 y Tabla 9.3; A9, métricas 9A-04 y siguientes; T13, carga.')
entry(s,'Compatibilidad y quince de quince integraciones','Las quince integraciones del alcance se prueban con sus flujos y contratos para verificar intercambio correcto y ausencia de errores.','No confundir quince integraciones con doce módulos ni con siete interfaces de usuario. Son inventarios de objetos distintos.','SD9 Tabla 9.3; T13 §§1, 3 y 8.2.')
entry(s,'Usabilidad WCAG 2 2 AA y guantes a menos veintidós grados','Reúne accesibilidad con uso real por los perfiles de terreno. El usuario debe poder operar el terminal en las condiciones físicas previstas, incluido el frío con guantes.','WCAG no demuestra por sí sola operación en cámara: la condición física requiere una prueba con personas y dispositivo.','SD9 Tabla 9.3 y §9.2.3; T13, perfil operacional.')
entry(s,'Fiabilidad y disponibilidad de al menos 99 coma 9 por ciento','El umbral mensual corresponde a la transacción crítica. El informe distingue además una disponibilidad de infraestructura de al menos 99,95 %.','No atribuir el mismo porcentaje a todos los componentes ni omitir la ventana protegida. Fiabilidad incluye continuidad y recuperación, no solo disponibilidad.','SD9 Tabla 9.3; A9, métricas de fiabilidad.')
entry(s,'RTO de hasta cuatro horas','RTO es el objetivo de tiempo para recuperar los servicios críticos ante el escenario de desastre definido. Se verifica con un ensayo de recuperación.','No es tiempo de respuesta de una pantalla ni permiso para interrumpir el despacho durante cuatro horas.','SD9 §§9.1.3 y 9.2.3; T13, recuperación ante desastres.')
entry(s,'RPO de hasta quince minutos','RPO es el objetivo máximo de pérdida temporal de datos que la recuperación admite en el escenario aplicable, medido respecto del estado recuperable.','No es duración del proceso de restauración ni plazo de sincronización de un terminal. Debe verificarse junto con el procedimiento de recuperación.','SD9 §9.1.3; A9, fiabilidad; T13, recuperación.')
entry(s,'Seguridad y cero hallazgos críticos o altos','Las pruebas y análisis no deben dejar abiertos hallazgos de seguridad de severidad crítica o alta al superar las puertas correspondientes.','No equivale a declarar riesgo cero o vulnerabilidades imposibles. Se requiere reevaluar al cambiar el software.','SD9 Tabla 9.3 y §9.2.2; A9, Anexo 9.B.')
entry(s,'Mantenibilidad y cobertura de setenta por ciento de negocio y ochenta global','Cobertura es la proporción de código ejercitada por las pruebas unitarias según la medición adoptada. El 70 % de lógica de negocio responde a la exigencia de las Bases; el 80 % global es compromiso adicional de LafroX.','No equivale a porcentaje de requisitos aceptados ni a probabilidad de que el programa sea correcto. Ambos umbrales se verifican y bloquean G2.','SD9 Tabla 9.4; A9, métricas 9A-30 y 9A-31; T13, pruebas unitarias.')
entry(s,'Portabilidad y mismo artefacto de QA a Producción','Se promueve la misma imagen o artefacto construido y verificado, identificado de forma inequívoca; no se recompila una versión distinta para Producción.','La configuración por ambiente puede variar bajo control. La afirmación no promete ejecutar en cualquier plataforma sin adaptación.','SD9 Tabla 9.3; A9, portabilidad; SD9 §9.2.5.')
s['script']=[
'La calidad se concreta en cuatro objetivos: probar todos los requisitos ofertados antes de su hito, cerrar las marchas blancas sin defectos críticos o altos, proteger la operación y verificar los dieciséis resultados del caso en su momento correspondiente.',
'Para medirlo usamos las ocho características de calidad. En desempeño no nos quedamos con promedios: el percentil 95 debe estar dentro del tiempo comprometido. Preparación tiene un segundo y preventa un segundo y medio bajo la carga definida.',
'En fiabilidad distinguimos la disponibilidad mensual de la transacción crítica de una exigencia más estricta: cero minutos de caída durante la ventana de despacho. Además, ensayamos la recuperación de servicios y datos con los objetivos de cuatro horas y quince minutos.',
'Probamos también las quince integraciones, la accesibilidad y el uso con guantes en frío. En mantenibilidad exigimos cobertura de pruebas y, para evitar diferencias entre lo probado y lo desplegado, promovemos el mismo artefacto.',
'Cada característica tiene una medida y un umbral. Ahora veremos qué ocurre cuando un resultado no cumple.'
]
s['notes']=['Si preguntan por p95, contrastarlo con promedio, no con disponibilidad.','RTO, RPO, disponibilidad y ventana protegida son cuatro conceptos diferentes.']

s=slide(36,'Calidad que bloquea','Idea central: las puertas técnicas impiden avanzar sin evidencia; la aceptación contractual pertenece al cliente y exige seis condiciones simultáneas.')
entry(s,'Siete puertas sin excepciones','Las puertas G0 a G6 son puntos de control de la cadena de entrega. Un incumplimiento bloqueante obliga a corregir y repetir la verificación correspondiente.','Sin excepciones no significa que todo cambio repita todas las campañas. G5 habilita cada etapa; las correcciones de marcha blanca recorren G0 a G4 y G6, con las pruebas aplicables.','SD9 §9.2.1, Tabla 9.7; A9, Anexo 9.B.')
entry(s,'G0 compromiso local formato y pruebas','Es el control sobre el cambio antes de incorporarlo al flujo compartido: formato del código y pruebas del módulo. Compromiso local traduce commit local.','No es un compromiso contractual ni aceptación de Puelche.','SD9 Figura 9.3; A9, Anexo 9.B, fila G0.')
entry(s,'G1 solicitud de fusión y revisión por pares','Una solicitud de fusión o merge request propone incorporar un cambio al código común. Requiere revisión y aprobación de otro integrante, junto con controles de secretos y estilo.','Revisar entre pares no reemplaza los análisis automáticos ni la aceptación del entregable por el cliente.','SD9 §9.2.2; A9, Anexo 9.B, fila G1.')
entry(s,'G2 integración continua cobertura complejidad y contratos','La cadena automática ejecuta pruebas y comprueba cobertura, complejidad y otras métricas del código, además de los contratos de interacción entre componentes.','Contrato aquí es una especificación técnica de intercambio, no el contrato comercial de la licitación. Una métrica bajo el umbral detiene la incorporación.','SD9 Tabla 9.4; A9, Anexo 9.B, fila G2.')
entry(s,'G3 promoción a QA sin vulnerabilidades críticas o altas','Promover es mover el artefacto al siguiente ambiente de control. QA significa aseguramiento de calidad; la puerta impide avanzar con vulnerabilidades bloqueantes y regresiones de desempeño.','Corregir primero; no registrar un hallazgo alto como deuda para avanzar igualmente.','SD9 §9.2.1; A9, Anexo 9.B, fila G3.')
entry(s,'G4 preproducción regresión DAST y WCAG','Regresión comprueba que el cambio no rompa capacidades existentes. DAST es análisis dinámico de seguridad del software ejecutándose; WCAG cubre accesibilidad.','La tabla de la lámina resume controles. El detalle de ambientes y ejecución lo fija T13; los análisis y las evidencias deben estar aprobados antes de promover.','SD9 §§9.2.1 a 9.2.3; A9, Anexo 9.B, fila G4.')
entry(s,'G5 entrada a Producción pruebas habilitantes y reversión','Se aplica al inicio de la marcha blanca de cada etapa, en H6 y H11. Exige H5 o H10 aceptado, pruebas aplicables, reversión ensayada, conciliación y capacitación.','No confundir la habilitación técnica de Producción con H7/H12, que aceptan el cierre contractual de marcha blanca.','SD9 §9.2.1; T17 §4.3; A9, Anexo 9.B.')
entry(s,'G6 producción supervisada canario y presupuesto de error','Canario es un despliegue gradual observado antes de extender el cambio. El presupuesto de error limita el margen de fallas compatible con el objetivo del servicio; agotarlo restringe despliegues no urgentes.','No autoriza indisponibilidad en la ventana protegida. El canario que empeora tiempos o errores activa la reversión definida.','A9, Anexo 9.B, fila G6; SD9 §9.2.1.')
entry(s,'Artículo 17 3 y seis condiciones para aceptar','Son copulativas: deben cumplirse todas para cerrar la marcha blanca. Ninguna compensa a otra. La aceptación se registra formalmente por el cliente.','No son las mismas siete puertas técnicas. LafroX verifica; la Contraparte Técnica acepta.','T17 §§1, 2 y 4.3, tabla de condiciones, p. 12.')
entry(s,'Cero incidentes críticos o altos','El protocolo exige cero incidentes críticos o altos abiertos atribuibles a la solución al cierre.','Se evalúan estado abierto y atribución. No afirmar que no ocurrió nunca un incidente durante todo el proyecto.','T17 §4.3, condición 1, p. 12.')
entry(s,'Cuatro semanas de volumen real','La condición completa es registrar en la solución el 100 % del volumen real de transacciones de la ola durante las cuatro últimas semanas.','La lámina abrevia una condición más exigente: no basta con usar una muestra ni con esperar cuatro semanas sin volumen efectivo.','T17 §4.3, condición 2, p. 12.')
entry(s,'Sin caída en despacho','Además de cero minutos entre 05:30 y 07:00, el protocolo verifica disponibilidad crítica y tiempos p95 de preparación, entrega, preventa y consulta de stock y crédito.','La frase abreviada no elimina los otros umbrales de disponibilidad y tiempo de respuesta.','T17 §4.3, condición 3, p. 12.')
entry(s,'Sin diferencias sin explicar','Conciliar exige que ninguna diferencia quede sin causa o explicación registrada. Puede haber diferencias identificadas y resueltas según el protocolo.','No reemplazarlo por nunca hay diferencias entre sistemas: la exigencia es cero diferencias no explicadas.','T17 §4.3, condición 4, p. 12.')
entry(s,'Cien por ciento de usuarios certificados','Se refiere a los usuarios de la ola que corresponde a la aceptación, capacitados y certificados con evidencia.','No confundir asistencia a una capacitación con certificación ni extender el denominador a cualquier persona ajena a esa ola.','T17 §4.3, condición 5, p. 12.')
entry(s,'Acta de la Contraparte','La Contraparte Técnica del cliente firma la aceptación con el expediente y los resultados exigidos. El protocolo contempla revisión y subsanación.','Un informe interno, el correo de un desarrollador o el despliegue por sí solos no sustituyen el acta.','T17 §§1, 2, 3 y 6; condición 6, p. 12.')
s['script']=[
'Para que los umbrales tengan efecto establecemos siete puertas, desde el cambio local hasta la producción supervisada. Las primeras revisan formato, pruebas y aprobación por pares. Luego se comprueban automáticamente cobertura, complejidad, contratos y seguridad.',
'Antes de promover un cambio verificamos regresión, seguridad dinámica y accesibilidad. Para habilitar cada etapa en Producción exigimos además certificación aceptada, conciliación de datos, capacitación y una reversión ensayada.',
'En producción observamos despliegues graduales. Si el cambio degrada el servicio, se aplica la reversión y se restringen nuevos despliegues según las reglas del plan.',
'Ahora bien, superar las puertas internas no significa aceptarnos a nosotros mismos. La aceptación corresponde a Puelche y exige las seis condiciones de la derecha: cero incidentes graves abiertos atribuibles a la solución; cien por ciento del volumen real durante las últimas cuatro semanas; disponibilidad y respuesta dentro de umbral; conciliación sin diferencias no explicadas; usuarios de la ola certificados y acta firmada.',
'Si falta una condición, el hito no se declara cerrado. Veamos cómo producimos la evidencia para tomar esas decisiones.'
]
s['notes']=['No confundir G5 con la aceptación contractual H7 o H12.','Al mencionar usuarios certificados, decir usuarios de la ola.']

s=slide(37,'Pruebas del código al terreno','Idea central: la estrategia combina pruebas automáticas con verificación humana y física; todas requieren resultados medibles o aceptación documentada.')
entry(s,'Doce tipos de prueba con número o firma','El plan agrupa doce tipos de prueba, desde unitarias hasta recuperación y despliegue. Número significa resultado medido; firma, validación o aceptación documentada por el responsable.','No implica que todo se automatice ni que cualquier firma reemplace un umbral técnico incumplido.','SD9 Tabla 9.8; T13 §3.')
entry(s,'Desarrollo unitarias componente y contratos de API','Las unitarias prueban lógica acotada; las de componente verifican unidades mayores; las de contrato comprueban que una API mantenga la estructura y comportamiento pactados. API es la interfaz de programación entre sistemas.','No sustituye una prueba de flujo completo con sistemas integrados.','T13 §3; SD9 §9.2.3.')
entry(s,'QA integración de quince flujos','Se ejecutan los intercambios de las quince integraciones del alcance y se verifica que no tengan errores.','Flujo aquí resume integración; no equivale a todos los escenarios posibles de negocio ni a quince módulos.','T13 §§1 y 3; SD9 Tabla 9.8.')
entry(s,'QA sistema y regresión','Sistema verifica funciones completas en conjunto. Regresión vuelve a probar lo existente después de un cambio para detectar efectos secundarios.','En la Etapa 2 también debe demostrarse que la Etapa 1 no se degradó.','SD9 Tabla 9.8; T13, criterios de salida; T17, H9/H12.')
entry(s,'QA DAST y accesibilidad','DAST busca fallas de seguridad sobre la aplicación en ejecución; accesibilidad combina herramientas con revisión manual.','El análisis dinámico no reemplaza análisis de código, dependencias, secretos o intrusión por un tercero.','SD9 §9.2.2; T13, pruebas de seguridad y accesibilidad.')
entry(s,'Preproducción aceptación de usuario','Usuarios de los perfiles relevantes prueban casos de negocio y la Contraparte Técnica deja evidencia de su aceptación.','Que el equipo técnico diga funciona no equivale a validar que satisface la necesidad del usuario.','SD9 Tabla 9.8 y §9.2.5; T13, aceptación de usuario.')
entry(s,'Carga de una vez y media y estrés','La carga verifica los umbrales a 1,5 veces el máximo de referencia. El estrés incrementa la demanda para observar crecimiento, punto de quiebre y degradación controlada.','No confundir 1,5 veces el máximo con un aumento del 150 %: es un 50 % adicional sobre la referencia.','SD9 §9.2.3 y Tabla 9.9; T13, carga y estrés.')
entry(s,'Resiliencia y sitios veinticuatro horas sin enlace','Se simulan fallas de infraestructura y dependencias; la propuesta incluye veinticuatro horas sin enlace en cada sitio, con continuidad y controles locales.','Es un escenario de sitio, distinto de las catorce horas de desconexión del terminal. No promete visibilidad central instantánea durante el corte.','SD9 Tabla 9.8; T13, resiliencia.')
entry(s,'Migración y despliegue','Migración verifica el traspaso y conciliación de datos. Despliegue comprueba la puesta en servicio y la reversión sin pérdida o interrupciones indebidas.','Reversión es volver a una versión o estado anterior controlado; no es restaurar arbitrariamente una base perdiendo datos nuevos.','SD9 Tabla 9.8; T13, migración y despliegue; T17, entrada a Producción.')
entry(s,'Intrusión por un tercero','Un especialista independiente realiza pruebas ofensivas autorizadas para detectar vulnerabilidades explotables antes de aceptar la entrada correspondiente.','Debe existir alcance y autorización; no equivale a ataques indiscriminados ni reemplaza la corrección de hallazgos.','SD9 Tabla 9.8; T13, pruebas de intrusión.')
entry(s,'Terreno catorce horas sin señal','Los terminales se someten a una desconexión prolongada y posterior sincronización para comprobar que no pierdan ni dupliquen operaciones.','No equiparar catorce horas sin señal con catorce horas de desactualización aceptable bajo cualquier regla; se comprueban además los criterios de sincronización.','SD9 Tabla 9.8; T13, perfil operacional y desconexión.')
entry(s,'Terreno menos veintidós grados con guantes','Comprueba que la interfaz y el dispositivo puedan usarse en cámara fría por la persona operadora con el equipo de protección previsto.','Es temperatura del escenario de uso, no una nueva meta universal de temperatura de todos los productos.','SD9 Tabla 9.8; T13, pruebas por perfil operacional.')
entry(s,'Terreno lluvia más de treinta y cuatro grados y una mano','Se verifica operación en condiciones físicas de ruta y trabajo: humedad, calor y manipulación limitada. El informe incorpora también sol en el perfil operacional.','No basta con una captura de pantalla de la aplicación: exige prueba física y criterios de salida del perfil.','SD9 Tabla 9.8; T13, perfil operacional.')
entry(s,'Producción recuperación RTO de hasta cuatro horas','Se ensaya el restablecimiento de los servicios críticos y se deja constancia del tiempo y estado recuperado. Debe verificarse también el RPO de quince minutos del plan.','Recuperación y resiliencia son distintas: una comprueba restaurar tras desastre; la otra, resistir o recuperarse de fallas específicas.','SD9 §§9.2.3 y 9.3.3; T13, recuperación y calendario.')
entry(s,'Marcha blanca e indicadores diarios','Durante operación supervisada se miden disponibilidad, tiempos, incidentes, conciliación y otras condiciones reales del servicio para sustentar el cierre.','Una prueba previa aprobada no sustituye la medición diaria con volumen real durante la marcha blanca.','SD9 §9.3.2; T17 §4.3.')
entry(s,'Mil trescientos dieciséis casos','Es el catálogo inicial dimensionado: 56 requisitos críticos × 5 casos + 163 altos × 3 + 43 medios × 2 + 356 RT × 1 + 15 integraciones × 7 = 1.316. RT significa requisito técnico.','Es una planificación inicial, no una cifra de pruebas ya ejecutadas. Las filas agrupadas y brechas se tratan como explica T13, sin duplicar casos.','T13 §8.2, pp. 19–20.')
entry(s,'Ochenta por ciento automatizados','Es el parámetro de diseño para automatizar unos 1.053 de los 1.316 casos; aproximadamente 263 requieren trabajo manual.','No es la cobertura global del código del 80 %: aquí el denominador son casos de prueba. Es una previsión a validar, no un resultado observado.','T13 §8.3, p. 20.')
entry(s,'Veintiuna coma noventa y nueve TPS','TPS significa transacciones por segundo. La referencia máxima es 14,66 TPS; multiplicada por 1,5 produce 21,99 TPS de carga de prueba.','No son personas simultáneas ni transacciones por minuto. La ventana protegida tiene además su carga y duración específicas.','SD9 Tabla 9.9; T13, dimensionamiento de carga.')
entry(s,'Cuatro coma cuatro horas y regresión en una noche','La estimación es 1.053 casos automáticos × 30 segundos ≈ 8,8 horas en un ejecutor; con dos ejecutores paralelos, ≈ 4,4 horas.','La frase de la lámina debe precisarse: corresponde a la porción automatizada. Los 263 casos manuales, a 15 minutos, requieren unas 65,8 HH por ciclo; no caben por definición en esas 4,4 horas.','T13 §8.3, p. 20; SD9 §9.2.3.')
entry(s,'Cero datos reales en pruebas sin anonimizar','Se utilizan datos sintéticos o datos reales anonimizados bajo el procedimiento del plan. Se evitan identidades y contenido sensible no autorizado en ambientes de prueba.','No decir que nunca se utiliza información derivada de datos reales; la condición es su tratamiento y anonimización antes del uso permitido.','SD9 §9.2.4; T13, datos de prueba.')
s['script']=[
'Las pruebas acompañan el recorrido completo: en desarrollo revisamos componentes y contratos; en QA, integración, sistema, seguridad y regresión. En preproducción verificamos aceptación, carga, resiliencia, migración y despliegue, incluyendo intrusión por un tercero.',
'La carga de referencia es de veintiuna coma noventa y nueve transacciones por segundo, una vez y media el máximo de septiembre utilizado en el dimensionamiento. Con estrés aumentamos la demanda para verificar que la solución controle la saturación sin perder transacciones.',
'Pero Puelche no opera en un laboratorio. Por eso verificamos terminales durante catorce horas sin señal, trabajo con guantes a menos veintidós grados y uso bajo lluvia, calor y con una mano. A nivel de sitio se prueba además continuidad durante veinticuatro horas sin enlace.',
'El catálogo inicial contempla mil trescientos dieciséis casos y un ochenta por ciento automatizado. Las cuatro coma cuatro horas corresponden solamente a esa parte automática, ejecutada en dos procesos paralelos. Las pruebas manuales y de terreno tienen programación propia.',
'Durante la marcha blanca medimos los indicadores diariamente con operación real. Y en pruebas usamos datos sintéticos o anonimizados. Todo resultado debe poder relacionarse con el requisito que demuestra; esa es la función de la trazabilidad.'
]
s['notes']=['Precisar siempre 4,4 horas de la parte automatizada, no de todo el trabajo manual y físico.','Distinguir 80 % de casos automatizados de 80 % de cobertura unitaria global.']

s=slide(38,'Trazabilidad y mejora','Idea central: vincular la necesidad con su implementación, prueba y aceptación permite demostrar cumplimiento y corregir causas, con esfuerzo de calidad programado.')
entry(s,'Retiro en menos de dos horas','El compromiso que abre el ejemplo es identificar a los clientes afectados por un lote en menos de dos horas mediante consulta y trazabilidad verificadas.','No promete retirar físicamente toda la mercadería en dos horas. La lámina usa una abreviación que debe precisarse al hablar.','SD9 §9.2.5; T17 §5, resultado de retiro sanitario; A9, Anexo 9.D.')
entry(s,'RF 09 01 trazabilidad','RF es requisito funcional. RF-09.01 identifica la capacidad de trazabilidad que se relaciona con el resultado de retiro.','Un identificador permite ubicar el requisito; no demuestra por sí solo que se haya implementado correctamente.','Diapositiva 38; SD9 Figura 9.7; A9, Anexo 9.D.')
entry(s,'M9 Calidad y trazabilidad','Es el módulo de la solución que concentra las capacidades del ejemplo de calidad y trazabilidad y participa en la consulta del lote y su evidencia.','No trabaja aislado: depende de registros de recepción, movimientos, entregas y datos relacionados.','Diapositiva 38; SD9 §9.2.5; A9, Anexo 9.D.')
entry(s,'Paquete EDT 3 4 4 cadena de frío y retiro','EDT es la estructura de desglose del trabajo. El código identifica la cuenta o agrupación de construcción que la lámina relaciona con esa capacidad.','Un vínculo EDT indica dónde se programa y controla el trabajo; no es el ID de una prueba ni una garantía de terminación.','Diapositiva 38; SD9 Figura 9.7; A9, Anexo 9.D y relación con T14.')
entry(s,'Simulacro de retiro','Prueba de negocio que parte de un lote y verifica la identificación de afectados, la evidencia y el tiempo comprometido.','No basta con probar que se abre una pantalla o se ejecuta una consulta: debe demostrarse el resultado del caso.','SD9 §§9.2.3 y 9.2.5; T13, aceptación; T17 §5.')
entry(s,'Fusión con el ID del requisito','La solicitud de fusión del código registra el identificador del requisito para poder seguir qué cambio lo implementó o corrigió.','No confundirlo con una unión de tablas ni con la firma de aceptación; es un enlace de trazabilidad del desarrollo.','SD9 §9.2.5 y Figura 9.7.')
entry(s,'Digest de la imagen firmada','Digest es la huella criptográfica que identifica exactamente la imagen de software. La firma permite verificar su procedencia e integridad en la cadena de entrega.','La huella identifica el artefacto; no demuestra por sí sola que el comportamiento sea correcto. Se vincula con los informes de la versión probada.','SD9 §9.2.5 y Figura 9.7; diapositiva 38.')
entry(s,'Acta del H7','H7 es el hito de paso contractual a producción de la Etapa 1, en el mes 16. La aceptación se respalda con resultados y condiciones de cierre de la marcha blanca.','No confundir H7 con la puerta G5 ni extender su acta a resultados cuyo momento de medición es posterior.','T17 §4.3 y §5; SD9 §9.2.5.')
entry(s,'PDCA y mejora continua','PDCA significa planificar, hacer, verificar y actuar. Se identifica causa y responsable, se aplica una corrección, se mide con el mismo criterio y se decide cerrar o escalar.','Aunque la lámina disponga Actuar y Verificar en posiciones gráficas distintas, el orden conceptual es Planificar → Hacer → Verificar → Actuar.','SD9 §9.2.6 y Figura 9.8.')
entry(s,'Diagramas de control','Permiten seguir una métrica en el tiempo y detectar señales de proceso fuera de control respecto de límites calculados.','Límite de control estadístico no es lo mismo que umbral contractual: estar estable no significa cumplir la exigencia.','SD9 §9.2.6.')
entry(s,'Pareto','Ordena causas o tipos de defecto por su contribución para priorizar los problemas que concentran más recurrencia o efecto.','No afirmar que siempre se cumple una proporción exacta 80/20; se utiliza la distribución observada.','SD9 §9.2.6.')
entry(s,'Causa raíz','Análisis de por qué se produce el defecto, más allá del síntoma. Una corrección se verifica y la prueba que detectó la falla se incorpora a la regresión.','Cerrar un ticket o repetir una prueba sin tratar la causa no constituye por sí solo mejora del proceso.','SD9 §9.2.6.')
entry(s,'Once mil ciento treinta y seis horas hombre','Suma programada de conformidad: 4.224 HH de prevención + 6.912 HH de evaluación = 11.136 HH. HH es esfuerzo humano acumulado, no duración calendario.','El título costo de la calidad es más amplio que esa cifra. El total integral debe considerar también esfuerzo por fallas, sin duplicar lo ya programado.','SD9 §9.1.6, Tabla 9.6; A9, Anexo 9.C.')
entry(s,'Cinco coma uno por ciento del total','Se calcula 11.136 / 217.543 × 100 ≈ 5,12 %, presentado redondeado como 5,1 %. El denominador es el esfuerzo total de la línea base.','No es porcentaje del precio económico ni tasa esperada de fallas.','SD9 §9.1.6; A9, Anexo 9.C.')
entry(s,'Prevención 37 coma 9 por ciento y evaluación 62 coma 1','Prevención busca evitar defectos; evaluación comprueba el cumplimiento con revisiones, certificaciones y pruebas. Las proporciones se calculan sobre las 11.136 HH de conformidad.','No son porcentajes de éxito ni prueban que evaluar sea intrínsecamente más importante; reflejan el esfuerzo programado del plan.','SD9 Tabla 9.6; A9, Anexo 9.C.')
s['script']=[
'Tomemos el compromiso de identificar a los clientes afectados por un lote en menos de dos horas. Ese resultado se vincula con un requisito, con el módulo M9, con el trabajo programado y con un simulacro que debe demostrarlo.',
'La cadena también conserva el cambio de código y la huella exacta de la imagen desplegada. Así, Puelche puede revisar qué se construyó, qué versión se probó y qué evidencia respalda su aceptación en el hito correspondiente.',
'Si una prueba falla, identificamos la causa y su responsable, corregimos, volvemos a medir y cerramos o escalamos la acción. Usamos diagramas de control y Pareto para distinguir problemas del proceso y priorizar las causas que más contribuyen.',
'Estas actividades tienen esfuerzo asignado: once mil ciento treinta y seis horas-hombre de prevención y evaluación, aproximadamente el cinco coma uno por ciento de la línea base. Es conformidad programada, no el costo integral de todas las fallas posibles.',
'Con esto cerramos calidad: cada compromiso tiene una medida, cada entrega tiene controles y cada aceptación tiene evidencia. A continuación veremos los riesgos que podrían impedir esos resultados y cómo responderemos.'
]
s['notes']=['Decir identificar afectados en menos de dos horas, no retirar físicamente todo.','No saltar directamente a la diapositiva 50: entre ambas está el bloque de riesgos.']

s=slide(50,'Cinco innovaciones','Idea central: cada innovación debe agregar una capacidad más allá del mínimo exigido y contar con indicador, validación y alternativa si no alcanza su resultado.')
entry(s,'Cinco innovaciones','La cartera cubre producto o servicio, proceso, tecnología o arquitectura, modelo de negocio o contratación e impacto social. T19 contiene una ficha por propuesta.','Son propuestas con madurez y validación pendientes, no cinco beneficios ya demostrados en Puelche.','S13, introducción y §§13.1 a 13.5; T19, fichas 1 a 5.')
entry(s,'Producto y vencimiento posentrega','INN-01 agrega seguimiento en el local después de entregar: estima saldo y ritmo de reposición, compara con vida útil y lleva al preventista avisos con acciones de un catálogo aprobado.','El saldo se confirma en la visita y no se convierte en inventario oficial del cliente. Es información estimada, no una medición continua de sus ventas.','S13 §§13.1.2 y 13.1.4; T19 ficha 1; A13, contrato aviso y respuesta de visita.')
entry(s,'Meta menos veinticinco por ciento de vencido','I-01B mide vencido en el local sobre entregado y compara con rutas de control sin avisos de los meses 13 a 15. La meta de reducción se mide en el mes 28.','No son 25 puntos porcentuales ni una reducción automática de toda la merma de inventario de Puelche. No sumar su beneficio al de INN-03 como si fueran independientes.','S13 §13.1.6; T19 ficha 1, campos 12 a 14; A13, Anexo 13.C.')
entry(s,'Proceso y reproducción de incidentes','INN-02 transforma un incidente en paquete sanitizado de evidencia, lo reproduce en QA o Preproducción y verifica allí la corrección; el escenario queda como regresión.','Agrega cierre por reproducción e incidentes no críticos a la base obligatoria. No abre un sexto ambiente ni replica datos personales sin control.','S13 §§13.2.2 y 13.2.4; T19 ficha 2; A13, paquete de reproducción.')
entry(s,'Meta menos treinta por ciento de horas de diagnóstico','I-02A busca reducir la mediana de horas hasta la causa reproducida, con al menos veinte pares comparables. La base son incidentes de prueba de QA de meses 5 a 10 diagnosticados sin el paquete; se mide en el mes 15.','No es promedio de horas, tiempo total de solución ni reducción prometida para todo incidente. La comparación empareja tipo de falla y módulo.','S13 §13.2.6; T19 ficha 2, campos 13 y 14; A13, Anexo 13.C.')
entry(s,'Tecnológica y vida útil por historia térmica','INN-03 integra temperatura y tiempo por lote mediante un modelo cinético determinista por familia, como Arrhenius o Q10, con parámetros aprobados. Apoya la decisión de Calidad.','El aporte es la integración operativa distribuida, no inventar esos modelos. No es un modelo de aprendizaje automático ni rellena tramos sin dato.','S13 §§13.3.2 y 13.3.4; T19 ficha 3; A13, contrato de historia térmica.')
entry(s,'Margen de más o menos quince por ciento','I-03B compara la estimación con la evaluación de Calidad. El margen se expresa respecto de la vida nominal del producto.','No equivale a un 15 % de ahorro, a un intervalo de confianza ni a quince días. La lámina abrevia el denominador que sí desarrolla el SD13.','S13 §§13.3.2 y 13.3.6; T19 ficha 3; A13, Anexo 13.C.')
entry(s,'En al menos ochenta por ciento de los lotes','El umbral de I-03B se comprueba por familia con al menos sesenta lotes por familia. La habilitación depende de que cada familia alcance la condición.','No agrupar todas las familias para ocultar una que falla. La estimación no extiende la fecha impresa ni desactiva bloqueos; sin datos suficientes se conserva la alternativa por fecha impresa.','S13 §13.3.6; T19 ficha 3, campos 13 a 17.')
entry(s,'Modelo de negocio y pago ligado al costo de servir','INN-04 agrega atribución y liquidación verificables de un tramo variable de Operación, con techo y base comparable. M10 calcula y Finanzas debe poder reproducirlo.','No presenta una tarifa nueva en la oferta técnica; los montos corresponden a la Oferta Económica. No se remunera cualquier variación externa como si fuera logro del proveedor.','S13 §§13.4.2 y 13.4.4; T19 ficha 4; A13, contrato de liquidación.')
entry(s,'Meta de al menos cinco por ciento menos costo','I-04A = 1 − costo comparable / costo de referencia. Usa una canasta e índices de ajuste definidos, con base de meses 21 a 23 y meta al mes 36; seguimiento mensual desde el 24.','No equivale a cinco puntos de margen comercial ni a ahorrar por excluir entregas difíciles del denominador. Primero hay tres liquidaciones en sombra; el efecto de pago comienza en el mes 24.','S13 §13.4.6; T19 ficha 4, campos 9 y 13 a 16.')
entry(s,'OTIF de al menos noventa y cinco por ciento','OTIF mide entregas a tiempo y completas. En INN-04 protege el servicio: la reducción de costo debe coexistir con el umbral de al menos 95 %.','No se considera éxito ahorrar degradando el servicio. La regla de atribución, los índices y el tratamiento de disputas se fijan contractualmente.','S13 §§13.4.2, 13.4.6 y 13.4.7; T19 ficha 4.')
entry(s,'Impacto social y hoja del almacenero','INN-05 devuelve al local información de compras, productos que dejó de pedir y productos por vencer. El propósito es que un comercio pequeño acceda a información útil sin necesitar sistema propio.','La comparación agregada es un cuarto bloque opcional condicionado a revisión legal; los bloques de datos propios pueden continuar si no se aprueba.','S13 §§13.5.2 y 13.5.4; T19 ficha 5; A13, contrato de hoja.')
entry(s,'Sin internet y pantalla','La hoja se calcula y precarga antes de la ruta. El preventista la muestra en su terminal sin exigir internet ni dispositivo al almacenero.','No es acceso ilimitado del cliente a un portal sin conexión: la atención se realiza mediante el terminal precargado del preventista.','S13 §13.5.4; T19 ficha 5, campos 4 y 7; A13, Anexo 13.A.')
entry(s,'Hoja impresa','La impresión utiliza la impresora de cabina ya ofertada y la copia se entrega si el local la solicita dentro del circuito previsto.','No prometer impresión inmediata durante toda visita: si no hay pedido, la copia espera otra entrega. Se mide por separado hoja mostrada, copia solicitada y copia entregada.','S13 §§13.5.4 y 13.5.6; T19 ficha 5; A13, indicadores I-05A y relacionados.')
s['script']=[
'Además de cumplir el alcance obligatorio y controlar sus riesgos, proponemos cinco innovaciones. Cada una agrega una capacidad concreta y tiene una forma de comprobar si aporta valor.',
'La primera lleva el seguimiento del vencimiento hasta el local del cliente. Con entregas, devoluciones y ritmo de reposición generamos avisos para que el preventista confirme la situación y aplique una acción aprobada. La meta es reducir un veinticinco por ciento el vencido en el local frente a la referencia definida, no toda la merma de Puelche.',
'La segunda transforma incidentes de terreno en escenarios reproducibles. En lugar de depender solo del relato de un pedido duplicado al recuperar señal, recogemos evidencia sanitizada, reproducimos la secuencia y validamos la corrección. La meta es reducir un treinta por ciento la mediana de horas hasta llegar a una causa reproducida.',
'La tercera usa la historia térmica para estimar vida útil remanente por lote. El modelo es determinista y se valida por familia. Buscamos que al menos el ochenta por ciento de los lotes quede dentro de un margen del quince por ciento de la vida nominal respecto de la evaluación de Calidad. Nunca amplía la fecha impresa ni elimina un bloqueo sanitario.',
'La cuarta alinea parte del pago de Operación con un resultado de Puelche. La meta es reducir al menos un cinco por ciento el costo comparable por entrega manteniendo un OTIF de al menos noventa y cinco por ciento. Primero se ensayan tres liquidaciones en sombra con Finanzas; no se considera éxito ahorrar deteriorando el servicio.',
'La quinta devuelve información de gestión al almacenero sin exigirle internet ni dispositivo. Puede revisar su hoja en la pantalla del preventista y solicitar una copia impresa para entrega posterior. La comparación con otros locales es opcional y depende de aprobación legal y controles de privacidad.',
'Las cinco propuestas aprovechan la arquitectura ofertada, pero requieren desarrollo y validación. Sus beneficios no se dan por demostrados: se miden y cada innovación tiene una alternativa si no alcanza su resultado.'
]
s['notes']=['En I-02A decir mediana, no promedio.','En INN-03 decir modelo determinista y precisión por familia; no inteligencia artificial que decide la inocuidad.','En INN-05 no prometer aumento de ventas: el efecto comercial se estima y valida.','Para pasar al cierre: Estas capacidades convergen en el compromiso con el que comenzamos, identificar afectados con evidencia y a tiempo.']

sources=[
('PPT','Presentación v10 alterna','Posiciones físicas 34, 35, 36, 37, 38 y 50.',PPT),
('SD9','Subdocumento 9 Plan de calidad','Objetivos, estándares, métricas, puertas, pruebas, mejora y programación.',BASE/'Subdocumento 9'/'LAFROX-Subdocumento9.pdf'),
('A9','Anexos del Subdocumento 9','9.A métricas; 9.B puertas; 9.C costo de calidad; 9.D trazabilidad de resultados.',BASE/'Subdocumento 9'/'LAFROX-Subdocumento9-Anexos.pdf'),
('T13','Formulario T 13','Plan de pruebas y validación; catálogo y cálculo de regresión en §8.',BASE/'Subdocumento 9'/'LAFROX-Formulario-T-13.pdf'),
('T17','Formulario T 17','Protocolo de aceptación por hito y de resultados; seis condiciones en p. 12.',BASE/'Subdocumento 9'/'LAFROX-Formulario-T-17.pdf'),
('S13','Subdocumento 13 Innovaciones','Diseño, alcance, madurez, indicadores y alternativas de las cinco propuestas.',BASE/'Subdocumento 13'/'LAFROX-Subdocumento13.pdf'),
('A13','Anexos del Subdocumento 13','13.A contratos de datos; 13.B trazabilidad y alcance; 13.C indicadores y riesgos.',BASE/'Subdocumento 13'/'LAFROX-Subdocumento13-Anexos.pdf'),
('T19','Formulario T 19','Fichas formales de las cinco innovaciones y condiciones de sus indicadores.',BASE/'Subdocumento 13'/'LAFROX-Formulario-T-19.pdf')
]

def verify_ppt():
    ns={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
    with zipfile.ZipFile(PPT) as z:
        p=ET.fromstring(z.read('ppt/presentation.xml'))
        rel=ET.fromstring(z.read('ppt/_rels/presentation.xml.rels'))
        relmap={r.attrib['Id']:r.attrib['Target'] for r in rel}
        ids=p.find('p:sldIdLst',ns)
        results={}
        for n in [34,35,36,37,38,50]:
            rid=ids[n-1].attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']
            target=relmap[rid]; part=target.lstrip('/') if target.startswith('/') else 'ppt/'+target
            x=ET.fromstring(z.read(part)); results[n]=[t.text or '' for t in x.findall('.//a:t',ns)]
        return len(ids),results

doc=Document()
sec=doc.sections[0]; sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=Inches(.75); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
sec.header_distance=Inches(.3); sec.footer_distance=Inches(.3)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st=doc.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.widow_control=True
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Normal'].paragraph_format.space_after=Pt(5)
doc.styles['Title'].font.size=Pt(26)
doc.styles['Title'].paragraph_format.space_after=Pt(12)
for name,size in [('Heading 1',20),('Heading 2',14),('Heading 3',11)]:
    st=doc.styles[name]; st.font.size=Pt(size); st.font.bold=True
    st.paragraph_format.space_before=Pt(12 if name!='Heading 3' else 8)
    st.paragraph_format.space_after=Pt(5)
    st.paragraph_format.keep_with_next=True
doc.styles['Subtitle'].font.size=Pt(12)
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
doc.core_properties.title='Diccionario de afirmaciones y guion de calidad e innovaciones'
doc.core_properties.subject='Diapositivas 34 35 36 37 38 y 50 de la presentacion v10 alterna'
doc.core_properties.author='LafroX'

def p(text='',boldlead=None,style=None):
    par=doc.add_paragraph(style=style)
    if boldlead:
        par.add_run(boldlead+' ').bold=True
    par.add_run(text)
    return par
def bookmark(par,name):
    idx=str(100+len(doc.paragraphs)); a=OxmlElement('w:bookmarkStart'); a.set(qn('w:id'),idx); a.set(qn('w:name'),name)
    b=OxmlElement('w:bookmarkEnd'); b.set(qn('w:id'),idx); par._p.insert(0,a); par._p.append(b)
def link(par,text,anchor):
    h=OxmlElement('w:hyperlink');h.set(qn('w:anchor'),anchor)
    r=OxmlElement('w:r');rp=OxmlElement('w:rPr');u=OxmlElement('w:u');u.set(qn('w:val'),'single');rp.append(u);r.append(rp)
    t=OxmlElement('w:t');t.text=text;r.append(t);h.append(r);par._p.append(h)

doc.add_paragraph('Diccionario de afirmaciones y guion de calidad e innovaciones',style='Title')
doc.add_paragraph('LafroX   Informe 2   Presentación v10 alterna',style='Subtitle')
p('Guía de estudio y exposición para las diapositivas 34, 35, 36, 37, 38 y 50. La numeración corresponde exclusivamente a la posición de cada diapositiva en PowerPoint, contando la portada como diapositiva 1.')
p('En cada diapositiva se presenta primero el diccionario de afirmaciones y, más abajo, el guion oral. El diccionario sirve para comprender, fundamentar y responder preguntas; el guion es una versión breve para exponer sin leer todas las definiciones.')
p('Una cifra presentada como meta o estimación no debe exponerse como resultado ya alcanzado. Las precisiones de alcance evitan ampliar los compromisos de la oferta.')
doc.add_heading('Cómo consultar la guía',level=2)
p('Cada entrada contiene la afirmación de la lámina, su significado, una precisión que limita interpretaciones incorrectas y el localizador de su fundamento. Los códigos de fuente se desarrollan al final del documento.')
p('Los rótulos de marca y navegación, como LafroX o Siguiente, no son afirmaciones técnicas y no forman parte del diccionario.')
doc.add_heading('Acceso por diapositiva',level=2)
for s in sections:
    par=p();link(par,f"Diapositiva {s['n']}   {s['title']}",f"slide{s['n']}")
p('El bloque de calidad ocupa las diapositivas 34 a 38. La diapositiva 50 de innovaciones se expone después del bloque de riesgos; no es la continuación inmediata de la 38.')
p('Tiempo orientativo del guion completo: entre 8 y 10 minutos, sujeto a ritmo y pausas. El diccionario no se lee durante la presentación.')

for s in sections:
    doc.add_page_break()
    h=doc.add_heading(f"Diapositiva {s['n']}",level=1);bookmark(h,f"slide{s['n']}")
    doc.add_heading(s['title'],level=2)
    p(s['intro'])
    doc.add_heading('Diccionario de afirmaciones',level=2)
    for i,(title,meaning,limit,source) in enumerate(s['entries'],1):
        doc.add_heading(f"{s['n']}.{i:02d}  {title}",level=3)
        p(meaning,boldlead='Significado').paragraph_format.keep_with_next=True
        p(limit,boldlead='Precisión').paragraph_format.keep_with_next=True
        pp=p(source,boldlead='Fundamento');pp.paragraph_format.space_after=Pt(8)
        for r in pp.runs:r.font.size=Pt(9)
    doc.add_page_break()
    doc.add_heading(f"Guion oral de la diapositiva {s['n']}",level=2)
    p('Texto para exponer. Las aclaraciones esenciales del diccionario ya están incorporadas al relato.')
    for txt in s['script']:p(txt)
    doc.add_heading('Recordatorios para la exposición',level=2)
    for txt in s['notes']:p(txt,style='List Bullet')

doc.add_page_break();doc.add_heading('Fuentes y localizadores',level=1)
p('Los localizadores de sección, tabla, figura y ficha remiten a los documentos siguientes. Cuando se indica página, corresponde a la página del PDF consultado. Las normas y referencias legales se explican en el contexto de la propuesta; no constituyen una auditoría jurídica independiente.')
for code,title,scope,path in sources:
    doc.add_heading(code+'  '+title,level=2)
    p(scope)
    pp=p(str(path));pp.paragraph_format.space_after=Pt(10)
    for r in pp.runs:r.font.size=Pt(8)

f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=f.add_run('LafroX   |   ');r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
count,texts=verify_ppt()
assert count==52, f'Cambió la presentación: {count} diapositivas'
assert 'Normas y leyes' in ' '.join(texts[34])
assert '5 innovaciones' in ' '.join(texts[50])
OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
audit={'slide_count':count,'sections':[{'slide':s['n'],'entries':len(s['entries']),'script_paragraphs':len(s['script'])} for s in sections], 'total_entries':sum(len(s['entries']) for s in sections),'source_slide_text':texts,'output':str(OUT)}
(ROOT/'tmp'/'guion_sd9_sd13'/'coverage.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in audit.items() if k!='source_slide_text'},ensure_ascii=False,indent=2))
