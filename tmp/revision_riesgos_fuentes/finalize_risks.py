from pathlib import Path
from docx import Document

path = Path("Datos externos/riesgos_conservados.docx")
doc = Document(path)

new_title_62 = "Llega un mensaje electrónico que aparenta venir de una cadena autorizada o fue alterado"
new_title_63 = "La integración recibe una posición falsa o antigua como si fuera actual"
for p in doc.paragraphs:
    if p.text.startswith("REC-062 — "):
        p.runs[0].text = "REC-062 — " + new_title_62
    if p.text.startswith("REC-063 — "):
        p.runs[0].text = "REC-063 — " + new_title_63
    if p.text.startswith("[Causa]") and "El seguimiento de la flota usa posiciones" in p.text:
        p.runs[3].text = " Al consultar los datos de seguimiento, podrían recibirse posiciones falsas, antiguas o fuera de orden presentadas como información actual.\n"
for row in doc.tables[1].rows[1:]:
    if row.cells[0].text == "REC-062":
        row.cells[1].text = new_title_62
    elif row.cells[0].text == "REC-063":
        row.cells[1].text = new_title_63

paras = doc.paragraphs
first = next(i for i,p in enumerate(paras) if p.text.startswith("REC-068 — "))
start = next(i for i,p in enumerate(paras) if p.text.startswith("REC-067 — "))
end = next(i for i,p in enumerate(paras) if p.text == "Distribución de los conservados en la RBS")
anchor = paras[first]._p
for p in paras[start:end]:
    anchor.addprevious(p._p)

doc.save(path)
print("ordered", path)
