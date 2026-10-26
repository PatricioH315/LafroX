from copy import deepcopy
from pathlib import Path
from docx import Document
import re

path = Path("Datos externos/riesgos_conservados.docx")
backup = Path("tmp/revision_riesgos_fuentes/riesgos_conservados_original.docx")
if not backup.exists():
    backup.write_bytes(path.read_bytes())

doc = Document(path)

edits = {
    "REC-012": {
        "title": "El Servicio de Impuestos Internos no responde y deja pendiente un trámite tributario",
        "cause": "La solución envía trámites tributarios por el sistema de gestión actual del cliente y necesita la respuesta del Servicio de Impuestos Internos (SII) para conocer su resultado.",
        "event": "Mientras se tramita uno de esos documentos, el SII podría no responder.",
        "effect": "El trámite quedaría sin confirmación. LafroX tendría que mantenerlo identificado como pendiente, seguir la incidencia y comprobar su resultado al reintentar la comunicación. Esto podría retrasar el cierre del proceso, pero la falta de respuesta no equivale a una aprobación.",
        "why": "LafroX debe dar seguimiento a los trámites que pasan por la integración y verificar su resultado antes de cerrar el proceso.",
        "contrast": "Ya está decidido que el sistema actual del cliente emite los documentos tributarios. El riesgo no cuestiona esa decisión: se refiere a una falta de respuesta del SII durante el trámite.",
    },
    "REC-014": {
        "title": "El proveedor de mensajes falla y un aviso operativo queda pendiente",
        "cause": "La solución utiliza servicios externos para enviar avisos operativos por los canales habilitados, como correo o mensajería.",
        "event": "Al enviar un aviso operativo, el proveedor del canal podría no responder o informar que el envío falló.",
        "effect": "El aviso quedaría pendiente o sin entrega comprobada. LafroX tendría que registrar el fallo, reintentar el envío según la política definida y verificar el resultado; esto podría retrasar la comunicación de la operación.",
        "why": "LafroX integra los canales de aviso y debe conservar evidencia del estado de cada envío. No puede garantizar que una persona lea el mensaje.",
        "contrast": "Los canales y la diferencia entre entrega y lectura ya están definidos. El riesgo es un fallo concreto del proveedor al enviar el aviso, no una decisión pendiente sobre qué canal usar.",
    },
    "REC-015": {
        "title": "El servicio externo de mapas limita las consultas y deja una planificación incompleta",
        "cause": "La solución consulta un servicio externo de mapas para mostrar ubicaciones y apoyar la planificación de rutas. Ese servicio aplica límites de uso.",
        "event": "Durante una consulta, el proveedor podría rechazar nuevas solicitudes porque se alcanzó su límite permitido.",
        "effect": "Las consultas afectadas no se completarían y la planificación que depende de ellas podría quedar pendiente. LafroX tendría que atender el incidente, comprobar cuándo se recupera el acceso y aplicar la continuidad ya prevista cuando corresponda.",
        "why": "LafroX mantiene la integración con el servicio de mapas y debe atender sus fallos; el efecto no se limita a una decisión de ruta del cliente.",
        "contrast": "La arquitectura ya contempla límites de uso y tiempos máximos de espera. El riesgo es que el proveedor aplique efectivamente una restricción durante el uso, no que LafroX aún deba decidir cómo integrar el servicio.",
    },
    "REC-035": {
        "title": "Un transportista no firma a tiempo el acuerdo necesario para incorporarlo al reparto",
        "cause": "Antes de incorporar a una empresa transportista al nuevo proceso de reparto, debe estar firmado su acuerdo operativo.",
        "event": "Una de las empresas transportistas podría no firmar el acuerdo antes de la ola de implantación en la que estaba prevista.",
        "effect": "Sus conductores no podrían incorporarse ni completar las pruebas en esa ola. LafroX tendría que reorganizar la capacitación y las pruebas, y quedaría pendiente demostrar la cobertura de ese transportista.",
        "why": "El acuerdo firmado permite que LafroX incorpore, capacite y pruebe a los conductores de esa empresa.",
        "contrast": "La obligación de firmar antes de incorporar al transportista ya está definida. El riesgo es la demora de una empresa concreta; dejarla fuera de una ola no elimina el alcance que debe acreditarse al cierre.",
    },
    "REC-036": {
        "title": "Un transportista no permite instalar el equipo que registra la temperatura de su camión",
        "cause": "La solución prevé instalar un termógrafo, equipo que registra la temperatura, en camiones refrigerados de terceros. Para hacerlo se necesita acceso al vehículo y autorización de su propietario.",
        "event": "Al preparar una ruta refrigerada, la empresa dueña del camión podría negar el acceso o la autorización para instalar ese equipo.",
        "effect": "LafroX no podría instalarlo ni integrarlo antes de incorporar esa ruta. La evidencia de temperatura que debía obtenerse mediante el equipo quedaría pendiente de una solución acordada y comprobada.",
        "why": "LafroX especifica e integra el equipo, aunque su compra corresponde al cliente. La negativa del transportista impide ejecutar y verificar esa parte de la implantación.",
        "contrast": "La responsabilidad de compra e integración ya está definida. Lo incierto es la autorización del dueño del camión; no se presupone que otro dato de temperatura cumpla automáticamente el requisito.",
    },
    "REC-042": {
        "title": "Los usuarios no pueden asistir a la capacitación prevista antes de cerrar una ola",
        "cause": "La capacitación y la certificación de usuarios deben realizarse sin detener la preparación de pedidos, la venta ni el reparto. Para ello, el cliente debe disponer de esas personas en los horarios acordados.",
        "event": "Las necesidades de la operación podrían impedir que algunos usuarios asistan a las sesiones de formación y evaluación antes del cierre de su ola de implantación.",
        "effect": "LafroX tendría que reprogramar las sesiones y el acompañamiento. No podría acreditar que esos usuarios están preparados para usar la solución hasta que completen la evaluación.",
        "why": "LafroX debe impartir la formación y presentar la evidencia de certificación de los usuarios de cada ola.",
        "contrast": "Las modalidades de capacitación sin detener la operación ya están definidas. El riesgo es que los usuarios no estén disponibles en las sesiones acordadas, no que falte decidir cómo capacitarlos.",
    },
    "REC-049": {
        "title": "Un corte de electricidad supera el tiempo de respaldo de equipos mantenidos por LafroX",
        "cause": "Algunos equipos de la infraestructura dependen del suministro eléctrico y cuentan con baterías o generación de respaldo cuya autonomía es limitada.",
        "event": "En un sitio atendido por LafroX, un corte de electricidad podría durar más que la autonomía del respaldo aplicable a los equipos afectados.",
        "effect": "Esos equipos podrían dejar de prestar servicio hasta que se restablezca la energía o se recupere la infraestructura. LafroX tendría que atender la incidencia y comprobar la restitución de los componentes bajo su mantenimiento.",
        "why": "El riesgo se limita a la sala técnica, los racks, los equipos de borde u otros componentes que LafroX mantiene; no atribuye a LafroX todas las pérdidas del negocio del cliente por un corte.",
        "contrast": "Los respaldos y sus autonomías ya están definidos. Lo incierto es que un corte externo dure más que el tiempo que esos respaldos pueden cubrir.",
    },
    "REC-057": {
        "title": "Falta la aprobación del cliente para tratar y respaldar datos en las regiones propuestas",
        "cause": "La solución propone operar y recuperarse usando regiones de nube en Brasil y Estados Unidos. El tratamiento y la transferencia internacional de los datos previstos requieren aprobación del cliente.",
        "event": "Cuando corresponda habilitar esos entornos, el cliente podría no haber otorgado la aprobación necesaria sobre la residencia y transferencia de los datos.",
        "effect": "LafroX no podría completar la habilitación o la evidencia de recuperación afectada hasta resolver formalmente esa condición con el cliente.",
        "why": "La aprobación es una condición para desplegar y demostrar los entornos que LafroX debe entregar.",
        "contrast": "Las regiones ya están elegidas en el diseño. El riesgo es no contar a tiempo con la aprobación requerida; la elección técnica no reemplaza esa autorización.",
    },
    "REC-061": {
        "title": "Una persona intenta usar la identidad de otra para hacer pedidos o entrar a un portal",
        "cause": "Los portales y otros accesos públicos permiten operar con credenciales o códigos de acceso personales, que un tercero podría intentar utilizar sin autorización.",
        "event": "Una persona podría presentar una credencial o un código ajeno para intentar registrar un pedido o acceder como si fuera su titular.",
        "effect": "LafroX tendría que detectar e investigar el intento, contener el uso indebido y comprobar qué operaciones fueron afectadas. Si el intento tuviera éxito, podrían quedar pedidos o accesos atribuidos a la persona equivocada.",
        "why": "La autenticación, los registros de actividad y la respuesta ante incidentes forman parte del servicio de LafroX.",
        "contrast": "Los controles de identidad ya están definidos y aquí no se afirma que hayan fallado. Se conserva el posible intento de suplantación y su efecto condicional.",
    },
    "REC-062": {
        "title": "Llega un pedido electrónico que aparenta venir de una cadena autorizada o fue alterado",
        "cause": "La solución recibe de cadenas comerciales mensajes electrónicos de pedidos y avisos de despacho a través de una conexión con sistemas externos.",
        "event": "Podría llegar un mensaje que aparenta provenir de una cadena autorizada sin serlo, o cuyo contenido haya sido cambiado durante el intercambio.",
        "effect": "LafroX tendría que separarlo de los mensajes válidos, comprobar su origen y contenido e investigar el incidente antes de procesar las operaciones afectadas. Si un mensaje falso se aceptara, podría alterar el pedido o aviso correspondiente.",
        "why": "LafroX implementa y opera la integración de intercambio electrónico y debe proteger la integridad de los mensajes que recibe.",
        "contrast": "Ya están definidos los certificados, las firmas y el rechazo de mensajes inválidos. El riesgo es recibir un mensaje hostil, no que aún falte decidir esos controles.",
    },
    "REC-063": {
        "title": "El proveedor de seguimiento entrega una posición falsa o antigua como si fuera actual",
        "cause": "El seguimiento de la flota usa posiciones enviadas por un proveedor externo. Para interpretarlas correctamente se necesita conocer su origen y la fecha de cada muestra.",
        "event": "El proveedor podría enviar una posición falsa, antigua o fuera de orden presentada como información actual.",
        "effect": "El seguimiento podría mostrar una ubicación incorrecta si la muestra no se detecta. LafroX tendría que comprobar su origen y fecha, apartar los datos afectados y resolver la incidencia de integración.",
        "why": "LafroX integra los datos de seguimiento y debe conservar la trazabilidad de las posiciones que incorpora a la solución.",
        "contrast": "El diseño ya define cómo ordenar y fechar las posiciones. El riesgo se refiere a los datos que efectivamente envíe el tercero, no a una decisión pendiente sobre el tratamiento de esas muestras.",
    },
}

headings = {}
for i, p in enumerate(doc.paragraphs):
    m = re.match(r"🟥?(REC-\d{3})\s+—", p.text)
    if m:
        headings[m.group(1)] = i

assert set(edits) <= set(headings), set(edits) - set(headings)

for risk, change in edits.items():
    i = headings[risk]
    heading = doc.paragraphs[i]
    heading.runs[0].text = risk + " — " + change["title"]
    description = doc.paragraphs[i + 5]
    assert description.text.startswith("[Causa]")
    description.runs[1].text = " " + change["cause"] + "\n"
    description.runs[3].text = " " + change["event"] + "\n"
    description.runs[5].text = " " + change["effect"]
    for offset, key, label in [(7, "why", "Por qué compete a LafroX:"), (9, "contrast", "Contraste con lo ya definido:")]:
        p = doc.paragraphs[i + offset]
        assert p.text.startswith(label), (risk, offset, p.text)
        p.runs[1].text = " " + change[key]

aliases = {
    "R-018": ("REC-068", "Reemplazos de conductores externos sin aviso", "Personas y adopción"),
    "R-027": ("REC-069", "Solicitudes que amplían el alcance", "Contrato y licitación"),
    "R-042": ("REC-070", "Cambios normativos", "Regulatorio y externo"),
}
for p in doc.paragraphs:
    for old, (new, title, branch) in aliases.items():
        if p.text.startswith(old + " — "):
            p.runs[0].text = new + " — " + title

doc.paragraphs[1].runs[-1].text = " De los 66 enunciados originales, 33 se conservan y 33 se excluyen. Con cuatro incorporaciones posteriores, esta lista contiene 37 riesgos."
doc.paragraphs[19].runs[0].text = (
    "Los 33 riesgos conservados del inventario original mantienen su código REC. "
    "Las cuatro incorporaciones posteriores reciben códigos REC-067 a REC-070; "
    "REC-068, REC-069 y REC-070 proceden respectivamente de R-018, R-027 y R-042, "
    "que se conservan como referencias de origen en sus fichas. "
    "Los saltos entre REC-001 y REC-066 corresponden a exclusiones; estos códigos no se equiparan automáticamente con el registro de SD8. "
    "Cada riesgo tiene una sola rama principal de la RBS existente."
)
doc.paragraphs[479].runs[0].text = "De los 66 originales, 33 se conservaron y 33 se excluyeron; se añadieron cuatro riesgos después, para un total de 37 fichas vigentes."
doc.paragraphs[480].runs[0].text = "Se verificaron 37 códigos REC únicos, descripciones completas, una rama por riesgo y la existencia de los archivos y páginas citados."

table = doc.tables[1]
table.rows[0].cells[0].text = "ID normalizada"
for risk, change in edits.items():
    matches = [row for row in table.rows if row.cells[0].text == risk]
    assert len(matches) == 1, risk
    matches[0].cells[1].text = change["title"]

for new, title, branch in [
    ("REC-067", "Accidente vehicular interrumpe un reparto incluido en pruebas o marcha blanca", "Operación del cliente"),
    *[(new, title, branch) for new, title, branch in aliases.values()],
]:
    row = deepcopy(table.rows[-1]._tr)
    table._tbl.append(row)
    for cell, value in zip(table.rows[-1].cells, (new, title, branch)):
        cell.text = value

assert len(table.rows) == 38
doc.save(path)
print(path, "updated", len(edits), "rewritten risks", len(table.rows) - 1, "total risks")
